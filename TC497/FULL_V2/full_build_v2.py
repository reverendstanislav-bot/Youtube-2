import os,csv,json,math,re,subprocess,wave
from pathlib import Path
from PIL import Image,ImageDraw,ImageFont,ImageEnhance,ImageFilter,ImageOps
from openpyxl import load_workbook
import numpy as np

PKG=Path(os.environ["PKG_ROOT"])
GEN=Path(os.environ["GEN_DIR"])
ARC=PKG/"04_DOWNLOADS"/"ARCHIVE_GREEN"
XLSX=PKG/"03_TIMELINE"/"TC497_FINAL_VISUAL_TIMELINE_CODEX_EDIT_MAP_v2.xlsx"
TR=Path(os.environ["TRANSCRIPT_DIR"])
EXTRA=Path(os.environ["V2_EXTRA"])
WORK=Path(os.environ["FULL_WORK"])
PUB=Path(os.environ["FULL_PUBLIC"])
IMGDIR=WORK/"state_images"; SEGDIR=WORK/"segments"; PDFDIR=WORK/"pdf_pages"
for p in [WORK,PUB,IMGDIR,SEGDIR,PDFDIR]: p.mkdir(parents=True,exist_ok=True)

FPS=30; CUT=82.6; END=1236.506; W,H=960,540
C={"charcoal":(23,26,28),"iron":(48,54,58),"paper":(230,221,200),"ivory":(243,235,221),"rust":(165,82,53),"blue":(95,116,125),"matte":(191,179,155)}
FONT_B="/usr/share/fonts/truetype/dejavu/DejaVuSansCondensed-Bold.ttf"
FONT_R="/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
def fnt(sz,b=True): return ImageFont.truetype(FONT_B if b else FONT_R,sz)
def sec(v): return float(v.total_seconds()) if hasattr(v,"total_seconds") else float(v)

CHAPTERS=[
("WHEN THE ROAD ENDS",82.6,200.699),
("THE ELECTRIC WHEEL",200.699,306.965),
("BIGGER AND BIGGER",306.965,399.491),
("THE 572-FOOT MACHINE",399.491,525.453),
("HOW DO YOU DRIVE 572 FEET?",525.453,663.014),
("NUCLEAR",663.014,750.184),
("YUMA",750.184,822.857),
("TEST RESULTS",822.857,915.853),
("IT WORKED",915.853,1002.057),
("DEFEATED BY THE SKY",1002.057,1115.376),
("THE LAST CAR",1115.376,1173.420),
("ENDING",1173.420,END)
]
EXTRA_BY_CH={
"WHEN THE ROAD ENDS":["V2_ROADLESS"],
"THE ELECTRIC WHEEL":["V2_ELECTRIC_DIST","V2_POWERED_HUB"],
"BIGGER AND BIGGER":["V2_PREDECESSOR"],
"HOW DO YOU DRIVE 572 FEET?":["V2_STEERING"],
"YUMA":["V2_YUMA_TEST"],
"TEST RESULTS":["V2_DUNE_LIMIT"],
"THE LAST CAR":["V2_LEGACY"],
"ENDING":["V2_LEGACY"]
}
EXTRA_FILES={
"V2_ROADLESS":"V2_ROADLESS.png",
"V2_ELECTRIC_DIST":"V2_ELECTRIC_DIST.png",
"V2_POWERED_HUB":"V2_POWERED_HUB.png",
"V2_PREDECESSOR":"V2_PREDECESSOR.png",
"V2_STEERING":"V2_STEERING.png",
"V2_YUMA_TEST":"V2_YUMA_TEST.png",
"V2_DUNE_LIMIT":"V2_DUNE_LIMIT.png",
"V2_LEGACY":"V2_LEGACY.png"
}
LIMITS={"AR-SNO":1,"AR-A01":2,"AR-OTTER2":4,"AR-OTTER1":2,"AR-DEW-MAP":2,"WB-L02":2,"WC-Y01":2,"WC-D01":2}
DEFAULT_LIMIT=2
BANNED_GRAPHICS={"GR-R01","GR-E01","GR-E02","GR-B01","GR-L01","GR-L02","GR-H01","GR-H02","GR-N01","GR-N02","GR-Y01","GR-T01","GR-T02","GR-W01"}

def fit_cover(im,size=(W,H),anchor=(.5,.5)):
    im=im.convert("RGB"); tw,th=size; iw,ih=im.size
    s=max(tw/iw,th/ih); nw,nh=round(iw*s),round(ih*s)
    im=im.resize((nw,nh),Image.Resampling.LANCZOS)
    x=max(0,min(nw-tw,round((nw-tw)*anchor[0]))); y=max(0,min(nh-th,round((nh-th)*anchor[1])))
    return im.crop((x,y,x+tw,y+th))
def fit_inside(im,size=(W,H),pad=22,bg=None):
    out=Image.new("RGB",size,bg or C["matte"]); im=im.convert("RGB")
    im.thumbnail((size[0]-2*pad,size[1]-2*pad),Image.Resampling.LANCZOS)
    out.paste(im,((size[0]-im.width)//2,(size[1]-im.height)//2)); return out
def grade(im,kind):
    im=im.convert("RGB")
    if kind=="generated":
        im=ImageEnhance.Brightness(im).enhance(1.035); im=ImageEnhance.Contrast(im).enhance(1.025); im=ImageEnhance.Color(im).enhance(.90)
    elif kind=="archive":
        im=ImageOps.grayscale(im).convert("RGB"); im=Image.blend(im,Image.new("RGB",im.size,C["paper"]),.10)
        im=ImageEnhance.Contrast(im).enhance(1.03); im=ImageEnhance.Brightness(im).enhance(.95)
    elif kind=="document":
        im=ImageEnhance.Color(im).enhance(.68); im=ImageEnhance.Brightness(im).enhance(.92); im=ImageEnhance.Contrast(im).enhance(1.0)
    return im

def gen_path(code):
    p=GEN/(code+".png")
    if not p.exists(): raise FileNotFoundError(code)
    return p
def archive_path(code):
    xs=list(ARC.glob(f"{code}__*"))
    if not xs: raise FileNotFoundError(code)
    return xs[0]
PDF_PAGES={"AR-OTTER1":1,"AR-OTTER2":1,"AR-PATENT":1}
def pdf_page(code,page=1):
    src=archive_path(code); out=PDFDIR/f"{code}_p{page}.png"
    if not out.exists():
        stem=out.with_suffix("")
        subprocess.run(["pdftoppm","-f",str(page),"-l",str(page),"-singlefile","-png","-r","130",str(src),str(stem)],check=True,stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
    return out

ANCHORS={"WA-09":(.52,.56),"WB-L05":(.46,.58),"WA-04R":(.54,.50),"WA-11R":(.52,.48),"WD-LC02":(.42,.50)}
def source_image(code,kind=None,page=None):
    if code in EXTRA_FILES:
        return grade(fit_cover(Image.open(EXTRA/EXTRA_FILES[code])), "generated")
    if kind=="Generated":
        return grade(fit_cover(Image.open(gen_path(code)),anchor=ANCHORS.get(code,(.5,.5))),"generated")
    if kind=="Archive":
        p=archive_path(code)
        if p.suffix.lower()==".pdf":
            pp=page or PDF_PAGES.get(code,1)
            return grade(fit_inside(Image.open(pdf_page(code,pp)),pad=24,bg=C["matte"]),"document")
        return grade(fit_inside(Image.open(p),pad=18,bg=C["matte"]),"archive")
    # resolve automatically
    gp=GEN/(code+".png")
    if gp.exists(): return source_image(code,"Generated")
    return source_image(code,"Archive",page)

def label(d,xy,text,size=18,color=None,b=True,anchor=None):
    d.text(xy,text,font=fnt(size,b),fill=color or C["ivory"],anchor=anchor)

def small_plate(d,text,x=36,y=34,color=None):
    bb=d.textbbox((0,0),text,font=fnt(15,True))
    w=bb[2]-bb[0]+24; h=bb[3]-bb[1]+14
    d.rounded_rectangle([x,y,x+w,y+h],6,fill=(23,26,28,205))
    label(d,(x+12,y+7),text,15,color or C["ivory"])

def overlay_semantic(im,semantic,chapter,shot_idx,local_idx):
    out=im.copy(); d=ImageDraw.Draw(out,"RGBA")
    # No production IDs, no full-frame cards, no decorative rust bars.
    if semantic in ("CP-R01","GR-R01"):
        # map markers only, never a route slash
        if semantic=="CP-R01":
            for x,y,n in [(214,235,"WEST"),(470,205,"DEW LINE"),(700,330,"SUPPLY GAP")]:
                d.ellipse([x-5,y-5,x+5,y+5],fill=C["blue"]+(235,))
                label(d,(x+11,y-9),n,13,C["ivory"])
        else:
            small_plate(d,"WHEN THE ROAD ENDS",36,36,C["paper"])
    elif semantic in ("GR-E01","CP-E01"):
        small_plate(d,"ELECTRIC DRIVE",36,36,C["paper"])
        # restrained evidence-first annotations
        for y,txt in [(390,"GENERATION"),(430,"ELECTRICAL DISTRIBUTION")]:
            d.line([(52,y),(150,y)],fill=C["blue"]+(220,),width=2)
            label(d,(162,y-10),txt,14,C["ivory"])
    elif semantic=="GR-E02":
        small_plate(d,"POWER AT THE WHEEL",36,36,C["paper"])
        d.line([(610,270),(760,230)],fill=C["blue"]+(220,),width=2)
        d.ellipse([602,263,616,277],fill=C["blue"]+(235,))
        label(d,(770,220),"MOTOR HOUSING",14,C["ivory"])
        d.line([(600,350),(755,382)],fill=C["blue"]+(220,),width=2)
        d.ellipse([593,343,607,357],fill=C["blue"]+(235,))
        label(d,(765,373),"POWERED HUB",14,C["ivory"])
    elif semantic in ("GR-B01","CP-B01"):
        small_plate(d,"OVERLAND TRAIN PREDECESSOR",36,36,C["paper"])
    elif semantic=="GR-L01":
        # measurement embedded on machine
        y=405; d.line([(100,y),(860,y)],fill=C["blue"]+(230,),width=3)
        for x in [100,290,480,670,860]: d.line([(x,y-9),(x,y+9)],fill=C["blue"]+(230,),width=2)
        label(d,(100,430),"572 FEET • 174 METERS",25,C["rust"])
    elif semantic in ("GR-L02","CP-L02"):
        small_plate(d,"13 UNITS • 54 DRIVEN WHEELS",36,36,C["paper"])
        label(d,(36,82),"DISTRIBUTED ELECTRIC DRIVE",14,C["ivory"])
    elif semantic in ("GR-H01","GR-H02","CP-H02"):
        # local steering traces, thin and functional
        pts1=[(120,390),(260,335),(430,330),(600,300),(815,270)]
        pts2=[(120,414),(260,365),(430,360),(600,330),(815,305)]
        d.line(pts1,fill=C["blue"]+(190,),width=2)
        d.line(pts2,fill=C["paper"]+(150,),width=2)
        small_plate(d,"CLOSER-FOLLOW STEERING",36,36,C["paper"])
    elif semantic=="CP-H01":
        small_plate(d,"PROJECT OTTER • STEERING EVIDENCE",36,36,C["paper"])
    elif semantic in ("GR-N01","CP-N01"):
        small_plate(d,"1961 NUCLEAR OVERLAND CONCEPT",36,36,C["paper"])
        label(d,(36,80),"CONCEPT — NOT THE TC-497 POWERPLANT",14,C["rust"])
    elif semantic in ("GR-N02","CP-N02"):
        small_plate(d,"TC-497 • GAS TURBINES + ELECTRIC DRIVE",36,36,C["paper"])
    elif semantic in ("GR-Y01","CP-Y01","CP-Y02"):
        small_plate(d,"YUMA PROVING GROUND",36,36,C["paper"])
        label(d,(36,82),"1962 → FEB 1963 → JUN 1963",15,C["ivory"])
    elif semantic in ("GR-T01","CP-T01"):
        small_plate(d,"PROJECT OTTER • TEST EVIDENCE",36,36,C["paper"])
    elif semantic in ("GR-T02","CP-T02"):
        small_plate(d,"DUNE LIMIT",36,36,C["paper"])
        label(d,(680,365),"≈12 FT",34,C["rust"])
        label(d,(698,405),"≈28°",24,C["ivory"])
    elif semantic in ("GR-W01","CP-W01"):
        metrics=[("~20 MPH","SPEED"),("~400 MI","RANGE"),("~150 TONS","CARGO"),("6","CREW")]
        n,l=metrics[local_idx%len(metrics)]
        d.rounded_rectangle([38,330,280,455],12,fill=(23,26,28,185))
        label(d,(58,345),n,43,C["rust"])
        label(d,(60,405),l,16,C["ivory"])
    elif semantic=="CP-D01":
        small_plate(d,"REAL U.S. ARMY CH-54",36,36,C["paper"])
    elif semantic=="CP-D02":
        small_plate(d,"THE LOGISTICS ANSWER CHANGED",36,36,C["paper"])
    return out.convert("RGB")

# canonical semantic rows
wb=load_workbook(XLSX,data_only=True); ws=wb["VISUAL_TIMELINE"]; headers=[c.value for c in ws[4]]
ROWS=[]
for r in ws.iter_rows(min_row=5,values_only=True):
    if not r[0]: continue
    d=dict(zip(headers,r)); d["Start"]=sec(d["Start"]); d["End"]=sec(d["End"]); d["Dur"]=sec(d["Dur"])
    ROWS.append(d)

def chapter_for(t):
    for c,a,b in CHAPTERS:
        if a<=t<b: return c,a,b
    return CHAPTERS[-1]
def row_for(t):
    for r in ROWS:
        if r["Start"]<=t<r["End"]: return r
    return ROWS[-1]

# word boundaries + punctuation
WORDS=[]
with open(TR/"TC497_VO_WORD_LEVEL.csv",encoding="utf-8-sig",newline="") as f:
    for r in csv.DictReader(f):
        s=float(r["start_sec"]); e=float(r["end_sec"])
        if s>=END: break
        WORDS.append({"s":s,"e":min(e,END),"w":r["word"]})

def phrase_boundaries(a,b):
    return [w["e"] for w in WORDS if a+2.7<=w["e"]<=b and (re.search(r"[.!?;:]$",w["w"]) or re.search(r",$",w["w"]))]

def make_cuts(a,b):
    targets=[4.0,5.8,3.4,6.6,4.6,7.2,3.8,5.2,6.0,4.3]
    punct=phrase_boundaries(a,b); out=[a]; t=a; k=0
    while b-t>7.8:
        goal=t+targets[k%len(targets)]; k+=1
        candidates=[x for x in punct if t+3.0<=x<=t+7.6]
        if candidates:
            nxt=min(candidates,key=lambda x:abs(x-goal))
        else:
            ends=[w["e"] for w in WORDS if t+3.1<=w["e"]<=min(b,t+7.4)]
            nxt=min(ends,key=lambda x:abs(x-goal)) if ends else min(b,t+5.0)
        if nxt-t<2.9: nxt=min(b,t+3.3)
        out.append(round(nxt,3)); t=nxt
    if b-out[-1]<2.5 and len(out)>1: out[-1]=b
    else: out.append(b)
    # dedupe
    clean=[out[0]]
    for x in out[1:]:
        if x-clean[-1]>.35: clean.append(x)
    if clean[-1]<b: clean[-1]=b
    return clean

# build chapter pools from unique physical assets
pools={}
for ch,a,b in CHAPTERS:
    arr=[]
    for x in EXTRA_BY_CH.get(ch,[]): arr.append((x,"Extra"))
    for r in ROWS:
        if r["End"]<=a or r["Start"]>=b: continue
        typ=r["Source Type"]; code=r["Asset Code"]
        if typ in ("Generated","Archive") and code not in [x[0] for x in arr]:
            if code=="AR-SNO" and ch!="WHEN THE ROAD ENDS": continue
            if code=="AR-A01" and ch not in ("THE 572-FOOT MACHINE","IT WORKED"): continue
            arr.append((code,typ))
    pools[ch]=arr

reuse={}
def allowed(code):
    return reuse.get(code,0)<LIMITS.get(code,DEFAULT_LIMIT)
def pick_asset(ch,semantic,mid,seq_i):
    r=row_for(mid); typ=r["Source Type"]; code=r["Asset Code"]
    # special evidence choices
    specials={
      "CP-R01":("AR-DEW-MAP","Archive"),
      "GR-E01":("V2_ELECTRIC_DIST","Extra"),
      "GR-E02":("V2_POWERED_HUB","Extra"),
      "CP-E01":("AR-PATENT","Archive"),
      "GR-B01":("V2_PREDECESSOR","Extra"),
      "GR-H01":("V2_STEERING","Extra"),
      "GR-H02":("V2_STEERING","Extra"),
      "CP-H02":("V2_STEERING","Extra"),
      "GR-Y01":("V2_YUMA_TEST","Extra"),
      "GR-T02":("V2_DUNE_LIMIT","Extra"),
      "CP-T02":("V2_DUNE_LIMIT","Extra")
    }
    if semantic in specials and allowed(specials[semantic][0]): return specials[semantic]
    if typ in ("Generated","Archive") and allowed(code):
        if code=="AR-SNO" and ch!="WHEN THE ROAD ENDS": pass
        elif code=="AR-A01" and ch not in ("THE 572-FOOT MACHINE","IT WORKED"): pass
        else: return code,typ
    pool=pools[ch]
    # rotate for diversity and respect caps
    for off in range(len(pool)):
        cand=pool[(seq_i+off)%len(pool)]
        if allowed(cand[0]): return cand
    # fallback permits one extra use of chapter-specific generated, never Sno
    for cand in pool:
        if cand[0] not in ("AR-SNO","AR-A01"): return cand
    return ("ST-COLD","Generated")

def image_for(code,typ,semantic,ch,shot_i,local_i):
    page=None
    if code=="AR-OTTER2":
        if semantic in ("GR-T01","CP-T01","CP-W01"): page=26
        elif semantic in ("GR-T02","CP-T02"): page=47
    im=source_image(code, "Generated" if typ=="Extra" else typ, page)
    # CP-D01 gets a true hero composite, the only intentional composite grammar.
    if semantic=="CP-D01":
        try:
            bg=source_image("WC-D01","Generated").convert("RGBA")
            heli=Image.open(archive_path("AR-CH54-1")).convert("RGB")
            crop=heli.crop((0,160,heli.width,min(760,heli.height)))
            crop.thumbnail((430,220),Image.Resampling.LANCZOS)
            gray=ImageOps.grayscale(crop)
            mask=gray.point(lambda p:255 if p<198 else max(0,int((232-p)*7.5))).filter(ImageFilter.GaussianBlur(1.3))
            rgba=crop.convert("RGBA"); rgba.putalpha(mask); bg.alpha_composite(rgba,(480,65)); im=bg.convert("RGB")
        except Exception: pass
    return overlay_semantic(im,semantic,ch,shot_i,local_i)

# build phrase-driven shot list
shots=[]; shot_i=0
for ch,a,b in CHAPTERS:
    cuts=make_cuts(a,b); local=0
    for s,e in zip(cuts[:-1],cuts[1:]):
        mid=(s+e)/2; r=row_for(mid); semantic=r["Asset Code"] if r["Source Type"] in ("Graphic","Composite") else ""
        code,typ=pick_asset(ch,semantic,mid,local)
        reuse[code]=reuse.get(code,0)+1
        shots.append({"i":shot_i,"s":s,"e":e,"chapter":ch,"semantic":semantic,"asset":code,"type":typ,"row":int(r["#"])})
        shot_i+=1; local+=1

# enforce final hard caps for Sno/A01 by replacing overflow in order
seen={}
for sh in shots:
    code=sh["asset"]; seen[code]=seen.get(code,0)+1
    lim=LIMITS.get(code,DEFAULT_LIMIT)
    if seen[code]>lim:
        pool=pools[sh["chapter"]]
        for cand in pool:
            if cand[0] in ("AR-SNO","AR-A01"): continue
            total=sum(1 for x in shots if x["asset"]==cand[0])
            if total<LIMITS.get(cand[0],DEFAULT_LIMIT):
                sh["asset"],sh["type"]=cand; break

def render_segment(sh,out):
    dur=sh["e"]-sh["s"]; frames=max(2,round(dur*FPS))
    chshots=[x for x in shots if x["chapter"]==sh["chapter"]]
    local_i=chshots.index(sh)
    img=image_for(sh["asset"],sh["type"],sh["semantic"],sh["chapter"],sh["i"],local_i)
    jpg=IMGDIR/f"{sh['i']:03d}_{sh['asset']}.jpg"; img.save(jpg,quality=94,subsampling=0)
    # Most shots are locked. Directed micro-push only on selected generated shots.
    motion=(sh["type"] in ("Generated","Extra") and dur>=5.6 and sh["i"]%5==1 and sh["semantic"]=="")
    if motion:
        vf=("scale=1008:567:force_original_aspect_ratio=increase,crop=1008:567,"
            "zoompan=z='min(zoom+0.000045,1.012)':x='iw/2-(iw/zoom/2)':y='ih/2-(ih/zoom/2)':d=1:s=960x540:fps=30,format=yuv420p")
    else: vf="scale=960:540,format=yuv420p"
    subprocess.run(["ffmpeg","-y","-loglevel","error","-framerate",str(FPS),"-loop","1","-i",str(jpg),"-vf",vf,
      "-frames:v",str(frames),"-an","-c:v","libx264","-preset","veryfast","-crf","19","-pix_fmt","yuv420p",str(out)],check=True)

paths=[]
for sh in shots:
    p=SEGDIR/f"{sh['i']:03d}.mp4"; render_segment(sh,p); paths.append(p)
concat=WORK/"rest_concat.txt"; concat.write_text("".join(f"file '{p.resolve()}'\n" for p in paths))
subprocess.run(["ffmpeg","-y","-loglevel","error","-f","concat","-safe","0","-i",str(concat),"-c","copy",str(PUB/"rest_base.mp4")],check=True)
(PUB/"timeline.json").write_text(json.dumps(shots,ensure_ascii=False,indent=2),encoding="utf-8")

# captions
phr=[]; buf=[]
for w in WORDS:
    buf.append(w); span=buf[-1]["e"]-buf[0]["s"]; chars=sum(len(x["w"])+1 for x in buf)
    punct=bool(re.search(r"[.!?;:]$",w["w"])); comma=bool(re.search(r",$",w["w"]))
    if len(buf)>=7 or chars>=48 or span>=2.45 or (punct and len(buf)>=3) or (comma and len(buf)>=5):
        phr.append({"s":buf[0]["s"],"e":buf[-1]["e"]+.10,"words":buf}); buf=[]
if buf: phr.append({"s":buf[0]["s"],"e":buf[-1]["e"]+.10,"words":buf)
(PUB/"captions.json").write_text(json.dumps(phr,ensure_ascii=False),encoding="utf-8")

# chapter and sparse source/concept bugs
events={"chapters":[],"bugs":[]}
for ch,a,b in CHAPTERS:
    events["chapters"].append({"s":a+.12,"e":min(a+2.65,b),"text":ch})
# Sparse, not every archive shot.
for ch,a,b in CHAPTERS:
    cshots=[x for x in shots if x["chapter"]==ch]
    arc=next((x for x in cshots if x["type"]=="Archive"),None)
    if arc: events["bugs"].append({"s":arc["s"]+.10,"e":min(arc["s"]+2.0,arc["e"]),"text":"ARCHIVE"})
for sh in shots:
    if sh["chapter"]=="NUCLEAR" and sh["asset"] in ("WB-N01","WB-N02"):
        events["bugs"].append({"s":sh["s"]+.1,"e":min(sh["s"]+2.0,sh["e"]),"text":"CONCEPT"}); break
(PUB/"events.json").write_text(json.dumps(events,ensure_ascii=False,indent=2),encoding="utf-8")

# chapter-specific restrained industrial bed starting at CUT
SR=48000; DUR=END-CUT; N=int(DUR*SR); t=np.arange(N,dtype=np.float32)/SR; score=np.zeros(N,dtype=np.float32)
for idx,(ch,a,b) in enumerate(CHAPTERS):
    aa=max(a,CUT)-CUT; bb=b-CUT; f=[41,47,52,44,49,38,45,50,54,42,39,43][idx]; amp=[.011,.013,.012,.013,.014,.011,.012,.014,.015,.017,.010,.009][idx]
    m=(t>=aa)&(t<bb); tt=t[m]-aa; env=np.ones_like(tt); fadeN=min(len(tt),int(1.0*SR))
    if fadeN>2:
        env[:fadeN]*=np.linspace(0,1,fadeN,dtype=np.float32); env[-fadeN:]*=np.linspace(1,0,fadeN,dtype=np.float32)
    sig=.78*np.sin(2*np.pi*f*tt)+.16*np.sin(2*np.pi*(f*1.5)*tt+.6)+.06*np.sin(2*np.pi*(f*2)*tt+1.1)
    score[m]+=amp*env*sig.astype(np.float32)
pcm=(np.clip(score,-.5,.5)*32767).astype(np.int16)
with wave.open(str(WORK/"rest_score.wav"),"wb") as wf:
    wf.setnchannels(1); wf.setsampwidth(2); wf.setframerate(SR); wf.writeframes(pcm.tobytes())

# audit
with open(WORK/"V2_SHOTS.csv","w",newline="",encoding="utf-8") as f:
    w=csv.writer(f); w.writerow(["i","start","end","dur","chapter","semantic","asset","type","row"])
    for x in shots: w.writerow([x["i"],x["s"],x["e"],round(x["e"]-x["s"],3),x["chapter"],x["semantic"],x["asset"],x["type"],x["row"]])
with open(WORK/"V2_REUSE.csv","w",newline="",encoding="utf-8") as f:
    w=csv.writer(f); w.writerow(["asset","uses"])
    for k,v in sorted({k:sum(1 for x in shots if x["asset"]==k) for k in set(x["asset"] for x in shots)}.items(),key=lambda z:(-z[1],z[0])): w.writerow([k,v])

durs=[x["e"]-x["s"] for x in shots]
print("FULL_V2_PREP_DONE",len(shots),"min",min(durs),"median",sorted(durs)[len(durs)//2],"max",max(durs))
print("AR-SNO",sum(1 for x in shots if x["asset"]=="AR-SNO"),"AR-A01-post",sum(1 for x in shots if x["asset"]=="AR-A01"))
